import hashlib
import html
import math
import re
import unicodedata

from openpyxl import load_workbook


PALETA_MAPA = {
    "ak": "#4D9DE0",
    "z": "#F2C230",
    "s1": "#48B96B",
    "sp": "#F29E38",
    "est": "#A66DD4",
    "s8": "#F06464",
    "s18": "#31B9BD",
    "ele": "#A8D800",
    "stock1": "#FF9F1C",
    "stock2": "#EF476F",
    "stock3": "#9B5DE5",
    "stock4": "#277DA1",
    "stock5": "#43AA8B",
    "stock6": "#F9C74F",
    "stock7": "#F15BB5",
    "stock8": "#00A6A6",
}

_PALETA_ADICIONAL = (
    "#1F77B4", "#FF7F0E", "#2CA02C", "#D62728", "#9467BD",
    "#8C564B", "#E377C2", "#7F7F7F", "#BCBD22", "#17BECF",
)
_ALIAS_GRUPOS = {
    "AK": "ak",
    "Z1Z11": "z",
    "S1S7": "s1",
    "SP1SP4": "sp",
    "ESTAB": "est",
    "S8S17": "s8",
    "S18S31": "s18",
    "ELE13": "ele",
}


def _normalizar(valor):
    texto = unicodedata.normalize("NFKD", str(valor or ""))
    texto = "".join(caracter for caracter in texto if not unicodedata.combining(caracter))
    return re.sub(r"[^A-Z0-9]", "", texto.upper())


def familia_seccion(nombre):
    clave = _normalizar(nombre)
    if clave in _ALIAS_GRUPOS:
        return _ALIAS_GRUPOS[clave]
    if re.fullmatch(r"[A-K]\d+", clave):
        return "ak"
    if re.fullmatch(r"Z(?:[1-9]|1[01])", clave):
        return "z"
    if re.fullmatch(r"S[1-7]", clave):
        return "s1"
    if re.fullmatch(r"SP[1-4]", clave):
        return "sp"
    if clave in ("ESTA", "ESTB"):
        return "est"
    if re.fullmatch(r"S(?:[89]|1[0-7])", clave):
        return "s8"
    if re.fullmatch(r"S(?:1[89]|2\d|3[01])", clave):
        return "s18"
    if re.fullmatch(r"ELE[1-3]", clave):
        return "ele"
    coincidencia = re.fullmatch(r"STOCK([1-8])", clave)
    if coincidencia:
        return f"stock{coincidencia.group(1)}"
    return None


def color_seccion(nombre, secciones_ordenadas=()):
    familia = familia_seccion(nombre)
    if familia:
        return PALETA_MAPA[familia]
    nombres = sorted({str(item) for item in secciones_ordenadas}, key=str.casefold)
    try:
        indice = next(i for i, item in enumerate(nombres) if item == str(nombre))
    except StopIteration:
        indice = int.from_bytes(hashlib.sha256(_normalizar(nombre).encode()).digest()[:2], "big")
    return _PALETA_ADICIONAL[indice % len(_PALETA_ADICIONAL)]


def _rango_celda(min_fila, min_columna, max_fila, max_columna, x, y):
    return (
        x[min_columna - 1],
        y[min_fila - 1],
        x[max_columna],
        y[max_fila],
    )


def _rango_marcador(ws, clave, celdas_marcador, x, y):
    celdas = celdas_marcador.get(clave, [])
    if not celdas:
        return []

    if re.fullmatch(r"S[1-7]", clave):
        return [
            _rango_celda(7, celda.column, 12, celda.column, x, y)
            for celda in celdas
        ]
    if re.fullmatch(r"S(?:1[89]|2\d|3[01])", clave):
        return [
            _rango_celda(celda.row, 16, celda.row, 23, x, y)
            for celda in celdas
        ]
    if clave in ("SP1", "SP2", "SP3", "SP4"):
        primera_fila, ultima_fila = (17, 23) if clave in ("SP1", "SP3") else (24, 28)
        return [
            _rango_celda(primera_fila, celda.column, ultima_fila, celda.column, x, y)
            for celda in celdas
        ]
    if clave in ("ESTA", "ESTB"):
        return [
            _rango_celda(17, celda.column, 28, celda.column, x, y)
            for celda in celdas
        ]
    if re.fullmatch(r"ELE[1-3]", clave):
        resultado = []
        for celda in celdas:
            for rango in ws.merged_cells.ranges:
                if rango.min_row == celda.row and rango.min_col == celda.column:
                    resultado.append(
                        _rango_celda(
                            rango.min_row, rango.min_col,
                            rango.max_row, rango.max_col, x, y,
                        )
                    )
                    break
        return resultado
    return [
        _rango_celda(celda.row, celda.column, celda.row, celda.column, x, y)
        for celda in celdas
    ]


def _rangos_grupo(ws, clave, celdas_marcador, x, y):
    if clave == "AK":
        celdas = [
            celda
            for nombre, celdas in celdas_marcador.items()
            if re.fullmatch(r"[A-K]\d+", nombre)
            for celda in celdas
        ]
        return [
            _rango_celda(celda.row, celda.column, celda.row, celda.column, x, y)
            for celda in celdas
        ]
    if clave == "Z1Z11":
        marcadores = [f"Z{numero}" for numero in range(1, 12)]
    elif clave == "S1S7":
        marcadores = [f"S{numero}" for numero in range(1, 8)]
    elif clave == "SP1SP4":
        marcadores = [f"SP{numero}" for numero in range(1, 5)]
    elif clave == "ESTAB":
        marcadores = ["ESTA", "ESTB"]
    elif clave == "S8S17":
        marcadores = [f"S{numero}" for numero in range(8, 18)]
    elif clave == "S18S31":
        marcadores = [f"S{numero}" for numero in range(18, 32)]
    elif clave == "ELE13":
        marcadores = [f"ELE{numero}" for numero in range(1, 4)]
    else:
        return []
    return [
        rango
        for marcador in marcadores
        for rango in _rango_marcador(ws, marcador, celdas_marcador, x, y)
    ]


def _subdividir_rangos(rangos, cantidad):
    if cantidad <= 0 or not rangos:
        return []
    resultado = []
    for indice, (x1, y1, x2, y2) in enumerate(rangos):
        numero = (cantidad + len(rangos) - indice - 1) // len(rangos)
        if numero <= 0:
            break
        ancho, alto = x2 - x1, y2 - y1
        columnas = max(1, min(numero, math.ceil(math.sqrt(numero * ancho / alto))))
        filas = math.ceil(numero / columnas)
        for posicion in range(numero):
            fila, columna = divmod(posicion, columnas)
            sx1 = x1 + ancho * columna / columnas
            sx2 = x1 + ancho * (columna + 1) / columnas
            sy1 = y1 + alto * fila / filas
            sy2 = y1 + alto * (fila + 1) / filas
            resultado.append((sx1, sy1, sx2, sy2))
            if len(resultado) == cantidad:
                return resultado
    return resultado


def _fmt_numero(valor):
    numero = float(valor or 0)
    return str(int(numero)) if numero.is_integer() else f"{numero:.1f}"


def _id_html(valor):
    return "detail-" + re.sub(r"[^a-z0-9]+", "-", _normalizar(valor).lower()).strip("-")


def crear_html_mapa(ruta_excel, secciones, ubicaciones_por_seccion, stock):
    ws = load_workbook(ruta_excel, data_only=True).active
    escala, margen, cabecera = 2, 56, 176

    anchos_columnas = []
    for indice in range(1, ws.max_column + 1):
        dim = ws.column_dimensions[ws.cell(1, indice).column_letter]
        caracteres = dim.width if dim.width is not None else (ws.sheet_format.defaultColWidth or 8.43)
        anchos_columnas.append(max(1, round((caracteres * 7 + 5) * escala)))
    alturas_filas = []
    for indice in range(1, ws.max_row + 1):
        dim = ws.row_dimensions[indice]
        puntos = dim.height if dim.height is not None else (ws.sheet_format.defaultRowHeight or 15)
        alturas_filas.append(max(1, round(puntos * 96 / 72 * escala)))

    x, y = [margen], [margen + cabecera]
    for ancho in anchos_columnas:
        x.append(x[-1] + ancho)
    for alto in alturas_filas:
        y.append(y[-1] + alto)
    ancho_svg, alto_svg = x[-1] + margen, y[-1] + margen

    celdas_marcador = {}
    for fila in ws.iter_rows():
        for celda in fila:
            if celda.value is not None:
                celdas_marcador.setdefault(_normalizar(celda.value), []).append(celda)

    inventario_por_ubicacion = {}
    for fila in stock:
        clave = (_normalizar(fila[2]), _normalizar(fila[1]))
        inventario_por_ubicacion.setdefault(clave, []).append(fila)

    datos_por_marcador = {}
    secciones_sin_mapa = []
    for seccion in secciones:
        seccion_id, nombre = seccion[0], str(seccion[1])
        ubicaciones = ubicaciones_por_seccion.get(seccion_id, [])
        candidatos = [_normalizar(nombre)]
        candidatos.extend(
            _normalizar(str(ubicacion[1]).split("-R", 1)[0])
            for ubicacion in ubicaciones
        )
        marcador = next(
            (
                candidato
                for candidato in candidatos
                if candidato in celdas_marcador or candidato in _ALIAS_GRUPOS
            ),
            None,
        )
        if marcador is None:
            secciones_sin_mapa.append(nombre)
            continue

        dato = datos_por_marcador.setdefault(
            marcador,
            {"nombre": nombre, "capacidad": 0, "ocupados": 0.0, "ubicaciones": []},
        )
        for ubicacion in ubicaciones:
            capacidad = max(0, int(ubicacion[4] or 0))
            ocupados = max(0.0, float(ubicacion[9] or 0))
            dato["capacidad"] += capacidad
            dato["ocupados"] += ocupados
            codigo = str(ubicacion[1])
            materiales = inventario_por_ubicacion.get(
                (_normalizar(nombre), _normalizar(codigo)),
                [],
            )
            dato["ubicaciones"].append(
                {
                    "codigo": codigo,
                    "capacidad": capacidad,
                    "ocupados": ocupados,
                    "materiales": materiales,
                }
            )

    svg = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{ancho_svg}" height="{alto_svg}" '
        f'viewBox="0 0 {ancho_svg} {alto_svg}" role="img" aria-label="Mapa interactivo del almacen Nave 1">'
        '<style>text{font-family:Arial,sans-serif}.slot:hover{stroke:#111820;stroke-width:5}'
        'a{cursor:pointer}.legend{font-size:24px;font-weight:bold}</style>'
        f'<rect width="{ancho_svg}" height="{alto_svg}" fill="#fff"/>'
        f'<text x="{margen}" y="50" font-size="50" font-weight="bold" fill="#202a36">'
        "MAPA DE OCUPACION - NAVE 1</text>"
        f'<text x="{margen}" y="82" font-size="27" fill="#657080">'
        'Color de zona igual a la grafica. Haz clic en una ubicacion para ver el contenido.</text>'
    ]

    leyenda = [
        ("A-K", "ak"), ("Z1-Z11", "z"), ("S1-S7", "s1"), ("SP1-SP4", "sp"),
        ("EST A-B", "est"), ("S8-S17", "s8"), ("S18-S31", "s18"), ("ELE 1-3", "ele"),
    ]
    leyenda.extend((f"STOCK {numero}", f"stock{numero}") for numero in range(1, 9))
    posicion_x, posicion_y = margen, 105
    for etiqueta, familia in leyenda:
        svg.append(
            f'<rect x="{posicion_x}" y="{posicion_y}" width="22" height="22" '
            f'rx="3" fill="{PALETA_MAPA[familia]}" stroke="#263238"/>'
            f'<text class="legend" x="{posicion_x + 29}" y="{posicion_y + 17}" fill="#273444">'
            f"{html.escape(etiqueta)}</text>"
        )
        posicion_x += 190
        if posicion_x > ancho_svg - 240:
            posicion_x, posicion_y = margen, posicion_y + 30

    # Draw occupied palette slots, leaving unused capacity white.
    for marcador, dato in datos_por_marcador.items():
        if marcador in _ALIAS_GRUPOS:
            rangos = _rangos_grupo(ws, marcador, celdas_marcador, x, y)
            etiqueta = dato["nombre"]
        else:
            rangos = _rango_marcador(ws, marcador, celdas_marcador, x, y)
            etiqueta = next(
                (str(celda.value) for celda in celdas_marcador.get(marcador, [])),
                dato["nombre"],
            )
        ranuras = _subdividir_rangos(rangos, dato["capacidad"])
        familia = familia_seccion(marcador) or familia_seccion(dato["nombre"])
        color = color_seccion(marcador, [dato["nombre"]])
        ocupados = dato["ocupados"]
        porcentaje = ocupados / dato["capacidad"] * 100 if dato["capacidad"] else 0
        detalle_id = _id_html(marcador)
        tooltip = [
            etiqueta,
            f'{_fmt_numero(ocupados)} / {dato["capacidad"]} palets ({porcentaje:.1f}%)',
        ]
        for ubicacion in dato["ubicaciones"]:
            if ubicacion["materiales"]:
                for material in ubicacion["materiales"]:
                    tooltip.append(
                        f"{material[3]}: {material[5]} palets, "
                        f"{material[7]} unidades sueltas"
                    )
            else:
                tooltip.append(f'{ubicacion["codigo"]}: sin materiales')
        title = html.escape(" | ".join(tooltip))
        for indice, (x1, y1, x2, y2) in enumerate(ranuras):
            # Capacity is distributed over the physical bays; extra spaces split a bay.
            capacidad_bahia = (dato["capacidad"] + len(ranuras) - indice - 1) // len(ranuras)
            if capacidad_bahia <= 0:
                break
            ocupacion_bahia = min(1.0, max(0.0, ocupados))
            ocupados -= ocupacion_bahia
            inset = 2
            bx1, by1 = x1 + inset, y1 + inset
            bw, bh = max(1, x2 - x1 - inset * 2), max(1, y2 - y1 - inset * 2)
            etiqueta_title = f'<title>{title}</title>'
            svg.append(
                f'<a href="#{detalle_id}" aria-label="{title}">{etiqueta_title}'
                f'<rect class="slot" x="{bx1:.2f}" y="{by1:.2f}" width="{bw:.2f}" '
                f'height="{bh:.2f}" rx="2" fill="#ffffff" stroke="{color}" stroke-width="2"/>'
            )
            if ocupacion_bahia:
                svg.append(
                    f'<rect pointer-events="none" x="{bx1 + 1:.2f}" y="{by1 + 1:.2f}" '
                    f'width="{max(0, (bw - 2) * ocupacion_bahia):.2f}" height="{max(0, bh - 2):.2f}" '
                    f'rx="1" fill="{color}"/>'
                )
            svg.append("</a>")

        if rangos:
            bx1 = min(rango[0] for rango in rangos)
            by1 = min(rango[1] for rango in rangos)
            bx2 = max(rango[2] for rango in rangos)
            by2 = max(rango[3] for rango in rangos)
            if bx2 - bx1 >= 180 or by2 - by1 >= 65:
                badge = f"{porcentaje:.0f}%"
                badge_width = max(34, len(badge) * 13)
                badge_x, badge_y = bx2 - badge_width - 2, by1 + 2
                svg.append(
                    f'<rect x="{badge_x:.2f}" y="{badge_y:.2f}" width="{badge_width}" height="26" '
                    'rx="4" fill="#ffffff" fill-opacity="0.93" stroke="#263238" stroke-width="1"/>'
                    f'<text x="{badge_x + badge_width / 2:.2f}" y="{badge_y + 19}" '
                    'font-size="18" font-weight="bold" text-anchor="middle" fill="#17212b">'
                    f"{badge}</text>"
                )

    # Restore the Excel floor-plan borders and labels above the occupancy overlay.
    for fila in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for celda in fila:
            izquierda, arriba = x[celda.column - 1], y[celda.row - 1]
            derecha, abajo = x[celda.column], y[celda.row]
            for lado, puntos in (
                ("left", (izquierda, arriba, izquierda, abajo)),
                ("right", (derecha, arriba, derecha, abajo)),
                ("top", (izquierda, arriba, derecha, arriba)),
                ("bottom", (izquierda, abajo, derecha, abajo)),
            ):
                borde = getattr(celda.border, lado)
                if borde.style:
                    ancho_linea = 4 if borde.style in ("medium", "thick", "double") else 2
                    color_linea = "#273444" if ancho_linea == 4 else "#718096"
                    svg.append(
                        f'<line x1="{puntos[0]}" y1="{puntos[1]}" x2="{puntos[2]}" '
                        f'y2="{puntos[3]}" stroke="{color_linea}" stroke-width="{ancho_linea}"/>'
                    )

    for celda in ws._cells.values():
        if celda.value is None:
            continue
        min_col = max_col = celda.column
        min_row = max_row = celda.row
        for rango in ws.merged_cells.ranges:
            if rango.min_col == celda.column and rango.min_row == celda.row:
                max_col, max_row = rango.max_col, rango.max_row
                break
        centro_x = (x[min_col - 1] + x[max_col]) / 2
        centro_y = (y[min_row - 1] + y[max_row]) / 2
        valor_texto = str(celda.value)
        ancho_texto = x[max_col] - x[min_col - 1] - 12
        tamano_fuente = max(18, min(36, int(ancho_texto / max(len(valor_texto) * 0.58, 1))))
        texto = html.escape(valor_texto)
        svg.append(
            f'<text x="{centro_x:.2f}" y="{centro_y:.2f}" font-size="{tamano_fuente}" '
            'font-weight="bold" text-anchor="middle" dominant-baseline="middle" fill="#17212b">'
            f"{texto}</text>"
        )
    svg.append("</svg>")

    detalles = ['<section class="details"><h3>Contenido y ocupacion por ubicacion</h3>']
    claves_ordenadas = sorted(
        datos_por_marcador,
        key=lambda clave: (
            familia_seccion(clave) or "",
            clave,
        ),
    )
    for marcador in claves_ordenadas:
        dato = datos_por_marcador[marcador]
        capacidad = dato["capacidad"]
        porcentaje = dato["ocupados"] / capacidad * 100 if capacidad else 0
        nombre = html.escape(dato["nombre"])
        detalles.append(
            f'<details id="{_id_html(marcador)}"><summary>{nombre}: '
            f'{_fmt_numero(dato["ocupados"])} / {capacidad} palets '
            f'({porcentaje:.1f}%)</summary>'
        )
        if not dato["ubicaciones"]:
            detalles.append("<p>Sin ubicaciones activas.</p>")
        for ubicacion in dato["ubicaciones"]:
            detalles.append(f'<p><strong>{html.escape(ubicacion["codigo"])}</strong>: ')
            detalles.append(
                f'{_fmt_numero(ubicacion["ocupados"])} / {ubicacion["capacidad"]} palets</p>'
            )
            if ubicacion["materiales"]:
                detalles.append("<ul>")
                for material in ubicacion["materiales"]:
                    nombre_material = html.escape(str(material[3]))
                    codigo_material = html.escape(str(material[4] or ""))
                    palets = int(material[5] or 0)
                    unidades_por_palet = int(material[6] or 1)
                    sueltas = int(material[7] or 0)
                    unidades = palets * unidades_por_palet + sueltas
                    sku = f" ({codigo_material})" if codigo_material else ""
                    detalles.append(
                        f"<li>{nombre_material}{sku}: {palets} palets, "
                        f"{sueltas} unidades sueltas ({unidades} unidades totales)</li>"
                    )
                detalles.append("</ul>")
            else:
                detalles.append("<p>Sin materiales registrados.</p>")
        detalles.append("</details>")
    detalles.append("</section>")

    return (
        '<!doctype html><html><head><meta charset="utf-8"><style>'
        'body{font-family:Arial,sans-serif;margin:0;color:#17212b}'
        'svg{display:block;width:100%;height:auto}'
        '.details{padding:12px 4px}.details h3{margin:0 0 10px}'
        'details{border:1px solid #cbd5e1;border-radius:6px;padding:8px 10px;margin:6px 0}'
        'summary{cursor:pointer;font-weight:600}.details p{margin:8px 0}'
        'details ul{margin:5px 0 8px;padding-left:24px}'
        '</style></head><body>'
        + "".join(svg)
        + "".join(detalles)
        + '<script>document.querySelectorAll(\'svg a[href^="#detail-"]\').forEach('
        'function(enlace){enlace.addEventListener("click",function(){'
        'var detalle=document.querySelector(enlace.getAttribute("href"));'
        'if(detalle){detalle.open=true;}});});</script>'
        + "</body></html>"
    ), secciones_sin_mapa
